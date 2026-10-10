from pathlib import Path
from typing import Any, Dict, List
from langchain_core.tools import tool
from docx import Document
from openpyxl import load_workbook


SUPPORTED_EXTENSIONS = {
    ".xlsx",
    ".docx",
    ".txt",
}


def detect_file_type(file_path: str) -> str:
    """
    Detect the type of a supported file.
    """

    extension = Path(file_path).suffix.lower()

    if extension == ".xlsx":
        return "excel"

    if extension == ".docx":
        return "word"

    if extension == ".txt":
        return "text"

    return "unknown"


def inspect_file_impl(file_path: str) -> Dict[str, Any]:
    """
    Inspect the structure of a file without extracting everything.
    """

    path = Path(file_path)

    if not path.exists():
        return {
            "status": "error",
            "error": f"File does not exist: {file_path}",
        }

    file_type = detect_file_type(file_path)

    if file_type == "excel":

        try:
            workbook = load_workbook(
                filename=file_path,
                read_only=True,
                data_only=False,
            )

            sheets = []

            for sheet_name in workbook.sheetnames:

                worksheet = workbook[sheet_name]

                rows = getattr(worksheet, "max_row", 0)
                columns = getattr(worksheet, "max_column", 0)

                if rows == 0 and columns == 0:
                    continue

                sheets.append(
                    {
                        "name": sheet_name,
                        "rows": rows,
                        "columns": columns,
                    }
                )

            workbook.close()

            return {
                "status": "success",
                "file": file_path,
                "file_type": "excel",
                "sheets": sheets,
            }

        except Exception as exc:
            return {
                "status": "error",
                "file": file_path,
                "error": str(exc),
            }

    if file_type == "word":

        try:
            document = Document(file_path)

            paragraphs = len(document.paragraphs)
            tables = len(document.tables)

            headings = []

            for paragraph in document.paragraphs:

                if paragraph.style and (paragraph.style.name or "").startswith("Heading"):
                    headings.append(paragraph.text)

            return {
                "status": "success",
                "file": file_path,
                "file_type": "word",
                "paragraphs": paragraphs,
                "tables": tables,
                "headings": headings,
            }

        except Exception as exc:
            return {
                "status": "error",
                "file": file_path,
                "error": str(exc),
            }

    if file_type == "text":

        try:

            with open(
                file_path,
                "r",
                encoding="utf-8",
            ) as file:

                content = file.read()

            return {
                "status": "success",
                "file": file_path,
                "file_type": "text",
                "characters": len(content),
                "lines": len(content.splitlines()),
            }

        except Exception as exc:
            return {
                "status": "error",
                "file": file_path,
                "error": str(exc),
            }

    return {
        "status": "error",
        "file": file_path,
        "error": f"Unsupported file type: {file_type}",
    }


def extract_text_impl(file_path: str) -> Dict[str, Any]:
    """
    Extract textual content from DOCX or TXT.
    """

    path = Path(file_path)

    if not path.exists():
        return {
            "status": "error",
            "error": f"File does not exist: {file_path}",
        }

    file_type = detect_file_type(file_path)

    if file_type == "word":

        try:

            document = Document(file_path)

            paragraphs = []

            for paragraph in document.paragraphs:

                text = paragraph.text.strip()

                if text:
                    paragraphs.append(text)

            return {
                "status": "success",
                "file": file_path,
                "file_type": "word",
                "paragraphs": paragraphs,
            }

        except Exception as exc:

            return {
                "status": "error",
                "file": file_path,
                "error": str(exc),
            }

    if file_type == "text":

        try:

            with open(
                file_path,
                "r",
                encoding="utf-8",
            ) as file:

                text = file.read()

            return {
                "status": "success",
                "file": file_path,
                "file_type": "text",
                "text": text,
            }

        except Exception as exc:

            return {
                "status": "error",
                "file": file_path,
                "error": str(exc),
            }

    return {
        "status": "error",
        "file": file_path,
        "error": "Text extraction is supported for DOCX and TXT files.",
    }


def extract_tables_impl(file_path: str) -> Dict[str, Any]:
    """
    Extract tables from DOCX or XLSX.
    """

    path = Path(file_path)

    if not path.exists():
        return {
            "status": "error",
            "error": f"File does not exist: {file_path}",
        }

    file_type = detect_file_type(file_path)

    if file_type == "word":

        try:

            document = Document(file_path)

            tables = []

            for table_index, table in enumerate(document.tables):

                rows = []

                for row in table.rows:

                    values = []

                    for cell in row.cells:
                        values.append(cell.text.strip())

                    rows.append(values)

                tables.append(
                    {
                        "table_index": table_index,
                        "rows": rows,
                    }
                )

            return {
                "status": "success",
                "file": file_path,
                "file_type": "word",
                "tables": tables,
            }

        except Exception as exc:

            return {
                "status": "error",
                "file": file_path,
                "error": str(exc),
            }

    if file_type == "excel":

        try:

            workbook = load_workbook(
                filename=file_path,
                read_only=True,
                data_only=True,
            )

            sheets = {}

            for sheet_name in workbook.sheetnames:

                worksheet = workbook[sheet_name]

                rows = []

                for row in worksheet.iter_rows(values_only=True):  
                    rows.append(list(row))

                sheets[sheet_name] = rows

            workbook.close()

            return {
                "status": "success",
                "file": file_path,
                "file_type": "excel",
                "sheets": sheets,
            }

        except Exception as exc:

            return {
                "status": "error",
                "file": file_path,
                "error": str(exc),
            }

    return {
        "status": "error",
        "file": file_path,
        "error": "Table extraction requires DOCX or XLSX.",
    }


def read_excel_range_impl(
    file_path: str,
    sheet_name: str,
    cell_range: str,
) -> Dict[str, Any]:
    """
    Read a specific Excel range.

    Example:
        A1:D20
    """

    path = Path(file_path)

    if not path.exists():

        return {
            "status": "error",
            "error": f"File does not exist: {file_path}",
        }

    try:

        workbook = load_workbook(
            filename=file_path,
            read_only=True,
            data_only=True,
        )

        if sheet_name not in workbook.sheetnames:

            workbook.close()

            return {
                "status": "error",
                "file": file_path,
                "error": (
                    f"Sheet '{sheet_name}' does not exist. "
                    f"Available sheets: {workbook.sheetnames}"
                ),
            }

        worksheet = workbook[sheet_name]

        rows = []

        for row in worksheet[cell_range]: 

            values = []

            for cell in row:
                values.append(cell.value)

            rows.append(values)

        workbook.close()

        return {
            "status": "success",
            "file": file_path,
            "file_type": "excel",
            "sheet": sheet_name,
            "range": cell_range,
            "rows": rows,
        }

    except Exception as exc:

        return {
            "status": "error",
            "file": file_path,
            "error": str(exc),
        }


def search_file_impl(
    file_path: str,
    query: str,
) -> Dict[str, Any]:
    """
    Search for a keyword/phrase inside Word, TXT, or Excel.
    """

    path = Path(file_path)

    if not path.exists():

        return {
            "status": "error",
            "error": f"File does not exist: {file_path}",
        }

    query_lower = query.lower()

    file_type = detect_file_type(file_path)

    matches = []

    if file_type == "word":

        document = Document(file_path)

        for index, paragraph in enumerate(
            document.paragraphs
        ):

            text = paragraph.text

            if query_lower in text.lower():

                matches.append(
                    {
                        "location": f"paragraph {index}",
                        "text": text,
                    }
                )

    elif file_type == "text":

        with open(
            file_path,
            "r",
            encoding="utf-8",
        ) as file:

            lines = file.readlines()

        for index, line in enumerate(lines):

            if query_lower in line.lower():

                matches.append(
                    {
                        "location": f"line {index + 1}",
                        "text": line.strip(),
                    }
                )

    elif file_type == "excel":

        workbook = load_workbook(
            filename=file_path,
            read_only=True,
            data_only=True,
        )

        for sheet_name in workbook.sheetnames:

            worksheet = workbook[sheet_name]

            for row in worksheet.iter_rows():

                for cell in row:

                    value = cell.value

                    if value is not None:

                        value_string = str(value)

                        if (
                            query_lower
                            in value_string.lower()
                        ):

                            matches.append(
                                {
                                    "sheet": sheet_name,
                                    "cell": cell.coordinate,
                                    "value": value,
                                }
                            )

        workbook.close()

    else:

        return {
            "status": "error",
            "file": file_path,
            "error": "Unsupported file type.",
        }

    return {
        "status": "success",
        "file": file_path,
        "query": query,
        "matches_found": len(matches),
        "matches": matches,
    }

import ast
import operator
from typing import Union

from langchain_core.tools import tool


# Allowed mathematical operators
OPERATORS = {
    ast.Add: operator.add,
    ast.Sub: operator.sub,
    ast.Mult: operator.mul,
    ast.Div: operator.truediv,
    ast.Pow: operator.pow,
    ast.Mod: operator.mod,
}


def evaluate_expression(node):
    """
    Safely evaluate a mathematical expression.
    Only numbers and basic arithmetic operators are allowed.
    """

    # Number
    if isinstance(node, ast.Constant):
        if isinstance(node.value, (int, float)):
            return node.value

        raise ValueError("Only numbers are allowed.")

    # Unary + or -
    if isinstance(node, ast.UnaryOp):
        value = evaluate_expression(node.operand)

        if isinstance(node.op, ast.USub):
            return -value

        if isinstance(node.op, ast.UAdd):
            return value

        raise ValueError("Unsupported unary operator.")

    # Binary operation
    if isinstance(node, ast.BinOp):

        left = evaluate_expression(node.left)
        right = evaluate_expression(node.right)

        operation = OPERATORS.get(type(node.op))

        if operation is None:
            raise ValueError("Unsupported mathematical operator.")

        if isinstance(node.op, ast.Div) and right == 0:
            raise ValueError("Division by zero.")

        return operation(left, right)

    raise ValueError(
        "Invalid expression. Only mathematical expressions are allowed."
    )



def calculator(expression: str) -> Union[int, float]:
    """
    General-purpose calculator.

    Evaluate a mathematical expression containing numbers
    and basic arithmetic operators.

    Supported:
        + addition
        - subtraction
        * multiplication
        / division
        % modulo
        ** power
        parentheses

    Examples:
        calculator("100 - 20 - 10 + 10")
        calculator("50 * 20")
        calculator("(100 + 50) * 0.14")
        calculator("500 + 500 * 0.14")

    The agent must determine the correct mathematical
    expression from the review process and policy.
    """

    try:
        tree = ast.parse(expression, mode="eval")

        result = evaluate_expression(tree.body)

        return result

    except Exception as e:
        raise ValueError(
            f"Could not calculate expression '{expression}': {e}"
        )


