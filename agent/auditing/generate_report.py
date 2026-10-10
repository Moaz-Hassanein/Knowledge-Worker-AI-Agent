from datetime import datetime
from pathlib import Path
from typing import Any, Dict

from docx import Document

from create_review_plan.state import GenericAgentState


def _write_report_docx(
    state: GenericAgentState,
    summary: Dict[str, Any],
    audit_entries: list,
    path: Path,
) -> None:

    doc = Document()

    doc.add_heading("Knowledge Worker Audit Report", level=1)

    doc.add_heading("Overview", level=2)
    doc.add_paragraph(
        f"Generated: {datetime.now():%Y-%m-%d %H:%M:%S}"
    )
    doc.add_paragraph(
        f"Use case: {state.get('use_case_description', '')}"
    )
    doc.add_paragraph(
        f"Source file: {state.get('target_file_path', '')}"
    )
    doc.add_paragraph(
        f"Status: {state.get('execution_status', 'UNKNOWN')}"
    )

    required_data = state.get("required_data", [])

    if required_data:
        doc.add_heading("Required Data", level=2)

        for field in required_data:
            doc.add_paragraph(f"- {field}")

    process_steps = state.get("process_steps", [])

    if process_steps:
        doc.add_heading("Approved Review Plan", level=2)

        for step in process_steps:
            doc.add_paragraph(
                f"{step['step_id']}. "
                f"{step['description']} "
                f"({step['action_type']})"
            )

    doc.add_heading("Summary", level=2)

    doc.add_paragraph(
        f"Total processed: {summary['total_processed']}"
    )

    doc.add_paragraph(
        f"Valid: {summary['valid']}"
    )

    doc.add_paragraph(
        f"Invalid: {summary['invalid']}"
    )

    doc.add_paragraph(
        f"Resolved by human: {summary['resolved_by_human']}"
    )

    doc.add_paragraph(
        f"- Ambiguous policy: {summary['ambiguous_count']}"
    )

    doc.add_paragraph(
        f"- Needed sign-off/escalation: "
        f"{summary['needs_intervention_count']}"
    )

    doc.add_heading("Audit Trail", level=2)

    table = doc.add_table(rows=1, cols=6)
    table.style = "Light Grid Accent 1"

    hdr = table.rows[0].cells

    hdr[0].text = "Record ID"
    hdr[1].text = "Status"
    hdr[2].text = "Resolved By"
    hdr[3].text = "Flagged Reason"
    hdr[4].text = "Reasoning"
    hdr[5].text = "Timestamp"

    flagged_labels = {
        "AMBIGUOUS": "Ambiguous policy",
        "NEEDS_INTERVENTION": "Needed sign-off/escalation",
    }

    for entry in audit_entries:

        row = table.add_row().cells

        row[0].text = str(entry.get("record_id", ""))
        row[1].text = str(entry.get("status", ""))

        row[2].text = (
            "Human"
            if entry.get("resolved_by_human")
            else "Agent"
        )

        row[3].text = flagged_labels.get(
            entry.get("flagged_reason"),
            "--"
        )

        row[4].text = str(
            entry.get("reasoning", "")
        )

        row[5].text = str(
            entry.get("timestamp", "")
        )

    if state.get("error_message"):

        doc.add_heading("Errors", level=2)

        doc.add_paragraph(
            str(state.get("error_message", ""))
        )

    path.parent.mkdir(parents=True, exist_ok=True)

    doc.save(str(path))


def generate_report(
    state: GenericAgentState,
) -> Dict[str, Any]:

    entries = state.get("audit_entries", [])

    # Build the summary HERE
    summary = {
        "total_processed": len(entries),

        "valid": sum(
            1
            for e in entries
            if e.get("status") == "VALID"
        ),

        "invalid": sum(
            1
            for e in entries
            if e.get("status") == "INVALID"
        ),

        "resolved_by_human": sum(
            1
            for e in entries
            if e.get("resolved_by_human")
        ),

        "ambiguous_count": sum(
            1
            for e in entries
            if e.get("flagged_reason") == "AMBIGUOUS"
        ),

        "needs_intervention_count": sum(
            1
            for e in entries
            if e.get("flagged_reason")
            == "NEEDS_INTERVENTION"
        ),
    }

    output_dir = Path("output")
    output_dir.mkdir(parents=True, exist_ok=True)

    report_path = (
        output_dir
        / f"review_report_{datetime.now():%Y%m%d_%H%M%S}.docx"
    )

    print("\n" + "=" * 70)
    print("GENERATE REPORT")
    print("=" * 70)

    print(f"Audit entries: {len(entries)}")
    print(f"Report path: {report_path}")

    # Generate DOCX
    _write_report_docx(
        state,
        summary,
        entries,
        report_path,
    )

    print(f"Report exists: {report_path.exists()}")

    error_message = state.get("error_message")

    status = (
        "REVIEW_COMPLETED"
        if not error_message
        else "FAILED"
    )

    return {
        "report_path": str(report_path),

        "execution_status": status,

        "is_completed": True,

        "final_result": {
            "report_path": str(report_path),
            "status": status,
            "summary": summary,
        },
    }